#!/usr/bin/env python3
"""
Iridium PowerPoint generator.

Usage:
    uv run --with python-pptx python3 create_powerpoint.py <content.json> [output.pptx]

content.json format:
{
  "title": "Presentation Title",
  "subtitle": "Optional subtitle",
  "slides": [
    {"type": "section",     "title": "Section Title"},
    {"type": "content",     "title": "Slide Title", "bullets": ["Point 1", "Point 2"]},
    {"type": "two_content", "title": "Slide Title", "left": ["Left 1"], "right": ["Right 1"]},
    {"type": "closing",     "title": "Questions?"}
  ]
}
"""
import json
import sys
from pptx import Presentation

TEMPLATE_PATH = '/home/alex99/Code/Internal/dotfiles/kiro/.kiro/templates/iridium_powerpoint_template.pptx'


def get_layout(presentation, layout_name):
    return next(layout for layout in presentation.slide_layouts if layout.name == layout_name)


def set_bullets(placeholder, bullet_points):
    text_frame = placeholder.text_frame
    text_frame.clear()
    for bullet in bullet_points:
        paragraph = text_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0


def remove_template_slides(presentation):
    while len(presentation.slides._sldIdLst):
        relationship_id = presentation.slides._sldIdLst[0].rId
        presentation.part.drop_rel(relationship_id)
        del presentation.slides._sldIdLst[0]


def add_cover_slide(presentation, title, subtitle):
    slide = presentation.slides.add_slide(get_layout(presentation, 'Custom Layout'))
    if slide.shapes.title:
        slide.shapes.title.text = title
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def add_section_slide(presentation, title):
    slide = presentation.slides.add_slide(get_layout(presentation, 'Section Header'))
    slide.shapes.title.text = title


def add_content_slide(presentation, title, bullets):
    slide = presentation.slides.add_slide(get_layout(presentation, 'Title and Content'))
    slide.shapes.title.text = title
    set_bullets(slide.placeholders[1], bullets)


def add_two_content_slide(presentation, title, left_bullets, right_bullets):
    slide = presentation.slides.add_slide(get_layout(presentation, 'Two Content'))
    slide.shapes.title.text = title
    set_bullets(slide.placeholders[1], left_bullets)
    set_bullets(slide.placeholders[2], right_bullets)


def add_closing_slide(presentation, title):
    slide = presentation.slides.add_slide(get_layout(presentation, 'Closing  (Option 1)'))
    if slide.shapes.title:
        slide.shapes.title.text = title


SLIDE_BUILDERS = {
    'section':     lambda presentation, slide_definition: add_section_slide(presentation, slide_definition['title']),
    'content':     lambda presentation, slide_definition: add_content_slide(presentation, slide_definition['title'], slide_definition.get('bullets', [])),
    'two_content': lambda presentation, slide_definition: add_two_content_slide(presentation, slide_definition['title'], slide_definition.get('left', []), slide_definition.get('right', [])),
    'closing':     lambda presentation, slide_definition: add_closing_slide(presentation, slide_definition.get('title', '')),
}


def build_presentation(content, output_path):
    presentation = Presentation(TEMPLATE_PATH)
    remove_template_slides(presentation)

    presentation.slides.add_slide(get_layout(presentation, 'Legal Disclosures'))
    add_cover_slide(presentation, content['title'], content.get('subtitle'))

    for slide_definition in content.get('slides', []):
        slide_type = slide_definition['type']
        builder = SLIDE_BUILDERS.get(slide_type)
        if builder:
            builder(presentation, slide_definition)
        else:
            print(f"Warning: unknown slide type '{slide_type}', skipping.")

    presentation.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    content = json.load(open(sys.argv[1]))
    output_path = sys.argv[2] if len(sys.argv) > 2 else 'presentation.pptx'
    build_presentation(content, output_path)
