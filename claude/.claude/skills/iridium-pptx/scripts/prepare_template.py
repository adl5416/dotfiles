"""Prepare a working copy of the Iridium PowerPoint template.

Copies the master template, strips the brand-guideline slides and the
"sample charts" note slide, and keeps only the sample slides for the
chosen theme (light or dark). Prints the remaining slides so the caller
knows the slide numbering of the working copy.
"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from pptx import Presentation

TEMPLATE_PATH = Path.home() / ".claude" / "templates" / "iridium_powerpoint_template.pptx"

# 1-based slide numbers in the original template.
GUIDELINE_SLIDES = {1, 2, 3}
CHART_NOTE_SLIDES = {12, 26}
LIGHT_SAMPLE_SLIDES = set(range(4, 18))
DARK_SAMPLE_SLIDES = set(range(18, 32))


def delete_slides(presentation: Presentation, slide_numbers: set[int]) -> None:
    """Delete slides from a presentation by 1-based slide number.

    Args:
        presentation: Open presentation to modify in place.
        slide_numbers: 1-based slide numbers to remove.
    """
    slide_id_list = presentation.slides._sldIdLst
    slide_ids = list(slide_id_list)
    for slide_number in sorted(slide_numbers, reverse=True):
        slide_id = slide_ids[slide_number - 1]
        presentation.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def main() -> None:
    """Copy the template, prune it to one theme, and report the result."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--theme", choices=("light", "dark"), default="light", help="Which theme's sample slides to keep")
    parser.add_argument("--output", type=Path, required=True, help="Path for the prepared .pptx working copy")
    parser.add_argument("--template", type=Path, default=TEMPLATE_PATH, help="Override the template location")
    arguments = parser.parse_args()

    shutil.copy(arguments.template, arguments.output)
    presentation = Presentation(arguments.output)

    other_theme_slides = DARK_SAMPLE_SLIDES if arguments.theme == "light" else LIGHT_SAMPLE_SLIDES
    delete_slides(presentation, GUIDELINE_SLIDES | CHART_NOTE_SLIDES | other_theme_slides)
    presentation.save(arguments.output)

    remaining_layouts = [slide.slide_layout.name for slide in Presentation(arguments.output).slides]
    print(f"Prepared {arguments.output}: {len(remaining_layouts)} {arguments.theme}-theme sample slides")
    for slide_number, layout_name in enumerate(remaining_layouts, start=1):
        print(f"  {slide_number}: {layout_name}")


if __name__ == "__main__":
    main()
