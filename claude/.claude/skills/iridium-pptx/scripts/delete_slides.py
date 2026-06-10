"""Delete slides from a .pptx file by 1-based slide number.

Use after deck content is complete to drop unused sample slides, e.g.:
    delete_slides.py deck.pptx 5 7 13
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation


def delete_slides(presentation: Presentation, slide_numbers: set[int]) -> None:
    """Delete slides from a presentation by 1-based slide number.

    Args:
        presentation: Open presentation to modify in place.
        slide_numbers: 1-based slide numbers to remove.
    """
    slide_id_list = presentation.slides._sldIdLst
    slide_ids = list(slide_id_list)
    for slide_number in sorted(slide_numbers, reverse=True):
        slide_id = slide_ids[slide_number - 1]
        presentation.part.drop_rel(slide_id.rId)
        slide_id_list.remove(slide_id)


def main() -> None:
    """Parse arguments and delete the requested slides."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx_path", type=Path, help="Presentation to modify in place")
    parser.add_argument("slide_numbers", type=int, nargs="+", help="1-based slide numbers to delete")
    arguments = parser.parse_args()

    presentation = Presentation(arguments.pptx_path)
    slide_count = len(presentation.slides)
    invalid_numbers = [number for number in arguments.slide_numbers if not 1 <= number <= slide_count]
    if invalid_numbers:
        parser.error(f"slide numbers out of range 1..{slide_count}: {invalid_numbers}")

    delete_slides(presentation, set(arguments.slide_numbers))
    presentation.save(arguments.pptx_path)
    print(f"Deleted {len(set(arguments.slide_numbers))} slide(s); {slide_count - len(set(arguments.slide_numbers))} remain in {arguments.pptx_path}")


if __name__ == "__main__":
    main()
